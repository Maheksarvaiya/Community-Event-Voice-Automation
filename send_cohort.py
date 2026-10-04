import os
import json
import requests
from datetime import datetime, date, timedelta
from dotenv import load_dotenv
from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

EXCEL_FILE = "community_event_automation.xlsx"

ORG_ID = "01a0c8c4-4468-735c-8ba8-45c5a3d42df0"
WORKSPACE_ID = "01a0c8c4-446f-784b-8bea-f5e77ae90ced"
CAMPAIGN_ID = "Community-E-9ad7c482-30dd"

# ------------------------------------------------------------
# SAFETY SWITCH
# ------------------------------------------------------------
# True  -> prepare cohort only, NO calls
# False -> actually submit cohort to Sarvam
#
# For normal production operation:
DRY_RUN = False


# ============================================================
# LOAD ENVIRONMENT
# ============================================================

load_dotenv()

SARVAM_API_KEY = os.getenv("VOICE_AGENTS_API_KEY")

if not SARVAM_API_KEY:
    raise RuntimeError(
        "VOICE_AGENTS_API_KEY was not found in .env"
    )


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def parse_excel_date(value):
    """
    Convert common Excel date formats into a Python date.
    """

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):

        value = value.strip()

        possible_formats = [
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
        ]

        for fmt in possible_formats:

            try:
                return datetime.strptime(
                    value,
                    fmt
                ).date()

            except ValueError:
                continue

    return None


def clean_phone(phone):
    """
    Normalize a phone number.
    """

    if phone is None:
        return None

    phone = str(phone).strip()

    # Remove spaces
    phone = phone.replace(" ", "")

    return phone


def valid_indian_phone(phone):
    """
    Basic validation for the E.164-style Indian number
    expected by this project.
    """

    if not phone:
        return False

    return (
        phone.startswith("+91")
        and len(phone) == 13
        and phone[1:].isdigit()
    )


def get_last_call_date(value):
    """
    Read Last Call from Excel.
    """

    if not value:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):

        value = value.strip()

        formats = [
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d",
        ]

        for fmt in formats:

            try:
                return datetime.strptime(
                    value,
                    fmt
                ).date()

            except ValueError:
                continue

    return None


# ============================================================
# LOAD EXCEL
# ============================================================

if not os.path.exists(EXCEL_FILE):
    raise FileNotFoundError(
        f"Excel file not found: {EXCEL_FILE}"
    )


workbook = load_workbook(EXCEL_FILE)

if "Events" not in workbook.sheetnames:
    raise RuntimeError(
        "Events sheet not found in Excel."
    )

if "Members" not in workbook.sheetnames:
    raise RuntimeError(
        "Members sheet not found in Excel."
    )


events_sheet = workbook["Events"]
members_sheet = workbook["Members"]


# ============================================================
# FIND TODAY / TOMORROW
# ============================================================

today = datetime.now().date()
tomorrow = today + timedelta(days=1)

print()
print("=" * 60)
print("COMMUNITY EVENT REMINDER AUTOMATION")
print("=" * 60)

print("Today:   ", today.strftime("%d-%m-%Y"))
print("Tomorrow:", tomorrow.strftime("%d-%m-%Y"))


# ============================================================
# FIND TOMORROW'S EVENTS
# ============================================================

tomorrow_events = []

for row_number, row in enumerate(
    events_sheet.iter_rows(
        min_row=2,
        values_only=True
    ),
    start=2
):

    event_id = row[0]
    event_name = row[1]
    event_date = row[2]
    event_time = row[3]
    location = row[4]

    if not event_id:
        continue

    event_date_only = parse_excel_date(event_date)

    if event_date_only is None:

        print(
            f"WARNING: Skipping event on Excel row "
            f"{row_number}: invalid date."
        )

        continue

    if event_date_only != tomorrow:
        continue

    if not event_name:
        print(
            f"WARNING: Skipping {event_id}: "
            "missing event name."
        )
        continue

    if not event_time:
        print(
            f"WARNING: Skipping {event_id}: "
            "missing event time."
        )
        continue

    if not location:
        print(
            f"WARNING: Skipping {event_id}: "
            "missing event location."
        )
        continue

    tomorrow_events.append({
        "event_id": str(event_id).strip(),
        "event_name": str(event_name).strip(),
        "event_date": event_date_only.strftime(
            "%d-%m-%Y"
        ),
        "event_time": str(event_time).strip(),
        "location": str(location).strip()
    })


# ============================================================
# NO EVENTS
# ============================================================

if not tomorrow_events:

    print()
    print("No events scheduled for tomorrow.")
    print("Nothing will be sent to Sarvam.")
    print("=" * 60)

    raise SystemExit(0)


# ============================================================
# DISPLAY EVENTS
# ============================================================

print()
print("-" * 60)
print("TOMORROW'S EVENTS")
print("-" * 60)

for event in tomorrow_events:

    print(
        f"{event['event_id']} | "
        f"{event['event_name']} | "
        f"{event['event_date']} | "
        f"{event['event_time']} | "
        f"{event['location']}"
    )


# ============================================================
# BUILD EVENT LOOKUP
# ============================================================

event_lookup = {
    event["event_id"]: event
    for event in tomorrow_events
}


# ============================================================
# FIND ELIGIBLE MEMBERS
# ============================================================

users = []

skipped_members = []

for row_number, row in enumerate(
    members_sheet.iter_rows(
        min_row=2,
        values_only=False
    ),
    start=2
):

    member_id = row[0].value
    customer_name = row[1].value
    phone = row[2].value
    member_event_id = row[3].value
    call_status = row[4].value
    last_call = row[6].value

    # --------------------------------------------------------
    # Basic validation
    # --------------------------------------------------------

    if not member_id:
        continue

    member_id = str(member_id).strip()

    if not customer_name:
        skipped_members.append(
            f"{member_id}: missing customer name"
        )
        continue

    if not member_event_id:
        skipped_members.append(
            f"{member_id}: missing event ID"
        )
        continue

    member_event_id = str(
        member_event_id
    ).strip()

    # Only members belonging to tomorrow's events
    if member_event_id not in event_lookup:
        continue

    if not phone:
        skipped_members.append(
            f"{member_id}: missing phone"
        )
        continue

    phone = clean_phone(phone)

    if not valid_indian_phone(phone):

        skipped_members.append(
            f"{member_id}: invalid phone"
        )

        continue

    # --------------------------------------------------------
    # DUPLICATE RUN PROTECTION
    # --------------------------------------------------------
    #
    # If this member was already processed today,
    # don't call them again if the script is accidentally
    # executed twice.
    #
    last_call_date = get_last_call_date(last_call)

    if last_call_date == today:

        print(
            f"SKIP {member_id}: "
            "already processed today."
        )

        continue

    # --------------------------------------------------------
    # BUILD USER
    # --------------------------------------------------------

    event = event_lookup[member_event_id]

    user = {
        "user_phone_number": phone,

        # IMPORTANT:
        # Member ID is the stable identity.
        "user_identifier": member_id,

        "app_variables": {
            "customer_name": str(customer_name).strip(),
            "community_name": "Our Community",
            "event_name": event["event_name"],
            "event_date": event["event_date"],
            "event_time": event["event_time"],
            "event_location": event["location"]
        }
    }

    users.append(user)


# ============================================================
# DISPLAY MEMBERS
# ============================================================

print()
print("-" * 60)
print("ELIGIBLE MEMBERS")
print("-" * 60)

if not users:

    print("No eligible members found.")
    print("No API request will be made.")
    print("=" * 60)

    raise SystemExit(0)


for user in users:

    print(
        f"{user['user_identifier']} | "
        f"{user['app_variables']['customer_name']} | "
        f"{user['app_variables']['event_name']} | "
        f"phone ending {user['user_phone_number'][-4:]}"
    )


print()
print("Total eligible members:", len(users))


# ============================================================
# SKIPPED MEMBERS SUMMARY
# ============================================================

if skipped_members:

    print()
    print("-" * 60)
    print("SKIPPED MEMBERS")
    print("-" * 60)

    for item in skipped_members:
        print("-", item)


# ============================================================
# BUILD COHORT NAME
# ============================================================

timestamp = datetime.now().strftime(
    "%Y%m%d-%H%M%S"
)

cohort_name = (
    f"community-reminder-"
    f"{tomorrow.strftime('%Y%m%d')}-"
    f"{timestamp}"
)


# ============================================================
# BUILD PAYLOAD
# ============================================================

payload = {
    "name": cohort_name,
    "users": users
}


# ============================================================
# DISPLAY COHORT SUMMARY
# ============================================================

print()
print("-" * 60)
print("COHORT SUMMARY")
print("-" * 60)

print("Cohort:", cohort_name)
print("Users:", len(users))
print("Dry run:", DRY_RUN)


# ============================================================
# DRY RUN
# ============================================================

if DRY_RUN:

    print()
    print("=" * 60)
    print("DRY RUN")
    print("=" * 60)

    print(
        json.dumps(
            payload,
            indent=4,
            ensure_ascii=False
        )
    )

    print()
    print("Nothing was sent to Sarvam.")
    print("No phone calls will be made.")
    print("=" * 60)

    raise SystemExit(0)


# ============================================================
# SARVAM API URL
# ============================================================

url = (
    "https://apps.sarvam.ai/api/scheduling/v1/"
    f"orgs/{ORG_ID}/"
    f"workspaces/{WORKSPACE_ID}/"
    f"campaigns/{CAMPAIGN_ID}/"
    "cohorts/stream"
)


# ============================================================
# API HEADERS
# ============================================================

headers = {
    "Content-Type": "application/json",
    "X-API-Key": SARVAM_API_KEY
}


# ============================================================
# SEND COHORT
# ============================================================

print()
print("=" * 60)
print("SUBMITTING COHORT TO SARVAM")
print("=" * 60)

try:

    response = requests.post(
        url,
        headers=headers,
        json=payload,
        timeout=30
    )

except requests.RequestException as error:

    print()
    print("ERROR: Could not connect to Sarvam.")
    print("Reason:", error)

    raise SystemExit(1)


print()
print("HTTP Status:", response.status_code)


# ============================================================
# PROCESS RESPONSE
# ============================================================

try:

    result = response.json()

except ValueError:

    result = {
        "raw_response": response.text
    }


print()
print(
    json.dumps(
        result,
        indent=4,
        ensure_ascii=False
    )
)


# ============================================================
# HANDLE API FAILURE
# ============================================================

if not response.ok:

    print()
    print("=" * 60)
    print("COHORT SUBMISSION FAILED")
    print("=" * 60)

    print(
        f"Sarvam returned HTTP {response.status_code}."
    )

    raise SystemExit(1)


# ============================================================
# SUCCESS
# ============================================================

print()
print("=" * 60)
print("COHORT SUBMITTED SUCCESSFULLY")
print("=" * 60)

print("Cohort:", cohort_name)
print("Members submitted:", len(users))

if isinstance(result, dict):

    cohort_id = result.get("cohort_id")

    if cohort_id:
        print("Cohort ID:", cohort_id)

print()
print(
    "Sarvam will process the cohort and "
    "send webhook results to your Flask webhook."
)

print("=" * 60)
