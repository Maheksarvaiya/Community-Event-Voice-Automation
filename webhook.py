from flask import Flask, request, jsonify
from openpyxl import load_workbook
from datetime import datetime

app = Flask(__name__)

EXCEL_FILE = "community_event_automation.xlsx"

ALLOWED_ATTENDANCE = {
    "CONFIRMED",
    "DECLINED",
    "MAYBE",
    "CALLBACK",
    "UNCLEAR"
}


def update_member_from_webhook(data):

    member_id = data.get("user_identifier")

    completion_status = data.get("completion_status")
    connectivity_status = data.get("connectivity_status")

    # Sarvam may send null when the call did not produce
    # agent output variables.
    output_variables = data.get("output_agent_variables") or {}

    attendance_response = output_variables.get(
        "attendance_response"
    )

    print("\n--- Webhook received ---")
    print("Member ID:", member_id)
    print("Completion:", completion_status)
    print("Connectivity:", connectivity_status)
    print("Attendance:", attendance_response)

    if not member_id:
        return False, "Missing user_identifier"

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Members"]

    for row in sheet.iter_rows(min_row=2):

        if str(row[0].value).strip() == str(member_id).strip():

            # ------------------------------------------
            # Determine call status
            # ------------------------------------------

            if connectivity_status == "connected":

                call_status = "COMPLETED"

            else:

                call_status = "FAILED"

            # ------------------------------------------
            # Validate attendance response
            # ------------------------------------------

            if attendance_response not in ALLOWED_ATTENDANCE:

                attendance_response = None

            # ------------------------------------------
            # Update Excel
            # ------------------------------------------

            row[4].value = call_status
            row[5].value = attendance_response
            row[6].value = datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

            workbook.save(EXCEL_FILE)

            print(
                f"Excel updated successfully for {member_id}"
            )

            return True, "Excel updated"

    return False, f"Member {member_id} not found"


@app.route("/webhook", methods=["POST"])
def webhook():

    data = request.get_json(silent=True)

    if not data:

        return jsonify({
            "success": False,
            "message": "Invalid JSON"
        }), 400

    success, message = update_member_from_webhook(data)

    if success:

        return jsonify({
            "success": True,
            "message": message
        }), 200

    return jsonify({
        "success": False,
        "message": message
    }), 400


if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )